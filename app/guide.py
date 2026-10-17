"""Life Group discussion guide for a sermon, written by the local LLM from the transcript.

Two passes so a 45-minute sermon fits a 7B model comfortably:
  1. notes: the sermon transcript in ~2,500-word chunks -> the points, verses, stories and lines
     the preacher actually used (grounding, so the questions stay accurate to the sermon);
  2. guide: those notes + the sermon details (+ 1-2 of your own past guides as examples of the
     format and tone, if GUIDE_EXAMPLES_PATH is set) -> numbered questions with leader notes.

Output is saved under data/guides/<job id>/ as guide.json, guide.md and a Word file.
"""
import html
import json
import logging
import random
import re
import threading
import time
from datetime import date
from pathlib import Path

from . import config, detect

log = logging.getLogger("sermon.guide")
DIR = config.DATA / "guides"
CHUNK_WORDS = 2500
# one context size for both passes, so Ollama doesn't reload the model in between: fits ~3,300 words of
# transcript (pass 1) or the notes of a long sermon + two example guides + the answer (pass 2)
NOTES_CTX = GUIDE_CTX = 12288
EXAMPLE_CHARS = 3500  # per example guide
_running = {}         # job id -> thread
_lock = threading.Lock()


# --- example guides (your own past guides, read at runtime; never stored by the app) -------------
def _date_key(p: Path):
    m = re.match(r"(20\d{6})", p.name)
    return m.group(1) if m else time.strftime("%Y%m%d", time.localtime(p.stat().st_mtime))


def _docx_text(path: Path):
    """A .docx as plain text, marking the parts the model should copy the shape of:
    numbered questions ("1."), their leader notes ("   - "), and other lines as they are."""
    import docx
    d = docx.Document(str(path))
    lines, qn = [], 0
    for p in d.paragraphs:
        t = " ".join(p.text.split())
        if not t:
            continue
        style = (p.style.name or "").lower()
        ppr = p._p.pPr
        num = ppr is not None and ppr.numPr is not None
        lvl = 0
        if num and ppr.numPr.ilvl is not None:
            lvl = int(ppr.numPr.ilvl.val or 0)
        if "answer" in style or (num and lvl >= 1):
            lines.append(f"   - {t}")
        elif "question" in style or num:
            qn += 1
            lines.append(f"{qn}. {t}")
        else:
            lines.append(t)
    return "\n".join(lines), qn


def examples(folder: Path | None, n=2):
    """Up to n of the newest guides in the folder (.docx, .md or .txt), as text."""
    if not folder or not folder.is_dir():
        return []
    files = []
    for p in folder.rglob("*"):
        if (p.suffix.lower() in (".docx", ".md", ".txt") and p.is_file() and not p.name.startswith(("~$", "."))
                and not any(".sb-" in part for part in p.parts)):
            files.append(p)
    out = []
    for p in sorted(files, key=_date_key, reverse=True):
        try:
            if p.suffix.lower() == ".docx":
                text, questions = _docx_text(p)
            else:
                text = p.read_text(errors="replace")
                questions = len(re.findall(r"(?m)^\s*\d+[.)]", text))
        except Exception:
            log.warning("guide: couldn't read example %s", p.name)
            continue
        if questions >= 4:  # a real discussion guide, not an empty template
            out.append(text[:EXAMPLE_CHARS])
        if len(out) >= n:
            break
    return out


# --- generation ----------------------------------------------------------------------------------
def _first_name(speaker):
    words = [w for w in re.split(r"\s+", speaker or "") if w and w.lower().rstrip(".") not in ("pastor", "pr", "rev", "dr", "brother")]
    return words[0] if words else ""


def _chunks(words, size=CHUNK_WORDS):
    if len(words) <= size * 1.3:
        return [words]
    n = round(len(words) / size)
    step = len(words) / n
    return [words[round(i * step):round((i + 1) * step)] for i in range(n)]


def _notes(text_words, info, status, stop):
    """Pass 1: what the preacher actually said, chunk by chunk."""
    parts = _chunks(text_words)
    notes = []
    for i, part in enumerate(parts):
        status(f"Reading the sermon ({i + 1} of {len(parts)})…")
        r = detect._ollama(
            f"This is part {i + 1} of {len(parts)} of a sermon transcript"
            + (f" on {info['scripture']}" if info.get("scripture") else "") + ". It was transcribed automatically, "
            "so some words may be wrong. Take short, accurate notes on THIS part for someone writing small-group "
            "discussion questions. Only write down what the preacher actually says here, not your own ideas:\n"
            "- points: up to 3 main points or headings he makes here, in order, each with a 1-2 sentence summary "
            "of what he said about it\n"
            "- verses: Bible references he reads or mentions (book chapter:verse)\n"
            "- stories: illustrations or stories he tells, one sentence each, saying who it is about (the "
            "preacher himself or someone else he tells about)\n"
            "- quotes: up to 2 memorable lines, nearly word for word\n"
            "- application: what he tells the listeners to believe or do (1-2 items)\n"
            "- questions: 2 discussion questions about the Bible text this part explains (specific words, verses, "
            "people or ideas: \"What does it mean that...\", \"Why do you think...\", \"What does verse ... show "
            "about...\"), each with the answer as the preacher explained it here (1-3 sentences, with verse numbers)\n"
            'Reply as JSON: {"points": [{"point": "...", "summary": "..."}], "verses": ["..."], '
            '"stories": ["..."], "quotes": ["..."], "application": ["..."], '
            '"questions": [{"q": "...", "answer": "..."}]}\n\n' + " ".join(part),
            NOTES_CTX, temperature=0.2, stop=stop)
        notes.append(r)
    return notes


def _notes_text(notes):
    out = []
    for i, n in enumerate(notes):
        out.append(f"Part {i + 1}:")
        for p in n.get("points") or []:
            if isinstance(p, dict):
                out.append(f"- POINT: {p.get('point', '')}: {p.get('summary', '')}")
            else:
                out.append(f"- POINT: {p}")
        for q in n.get("questions") or []:
            if isinstance(q, dict) and q.get("q"):
                out.append(f"- CANDIDATE QUESTION: {q['q']} | ANSWER: {q.get('answer', '')}")
        for key, label in (("verses", "VERSES"), ("stories", "STORY"), ("quotes", "QUOTE"), ("application", "APPLICATION")):
            vals = [str(v) for v in (n.get(key) or []) if str(v).strip()]
            if key == "verses" and vals:
                out.append(f"- VERSES: {', '.join(vals)}")
            else:
                out.extend(f"- {label}: {v}" for v in vals)
    return "\n".join(out)


GENERIC_FORMAT = """Life Group Questions
"Title"—Passage
1. A light, personal opening question.
   - Leader note.
2. A question about the passage or a point the preacher made.
   - The answer as the sermon gave it, with verse numbers.
3. ...
   - ..."""


def _prompt(info, notes_text, exs):
    who = _first_name(info.get("speaker")) or "the preacher"
    head = (f"Sermon: \"{info.get('title') or 'untitled'}\""
            + (f" on {info['scripture']}" if info.get("scripture") else "")
            + (f", part of the series \"{info['series']}\"" if info.get("series") else "")
            + (f", preached by {info['speaker']}" if info.get("speaker") else "") + ".")
    if exs:
        fmt = ("Here are discussion guides this leader wrote before. Match their TONE, LENGTH and the KIND of "
               "questions they ask (questions that dig into the meaning of the text, each followed by short leader "
               "notes with the answer). Do NOT copy their topics, wording or the leader's personal stories.\n\n"
               + "\n\n".join(f"=== Example guide {i + 1} ===\n{e}" for i, e in enumerate(exs)))
    else:
        fmt = "The guide looks like this:\n" + GENERIC_FORMAT
    return (
        "You write the weekly discussion guide for a church small group (Life Group) that meets during the week "
        "to talk about Sunday's sermon.\n\n" + fmt + "\n\n"
        f"=== The new sermon ===\n{head}\nNotes taken from the sermon transcript, in order:\n{notes_text}\n\n"
        "First work out the big idea of the sermon (one sentence) and its main points (3-5 short phrases of 3-8 "
        "words, in the order the preacher made them; teaching points, not \"Introduction\" or \"Historical context\"). Then write 7 questions in this order:\n"
        "1. OPENER: a light, personal icebreaker anyone can answer from their own life (a memory, a habit, an "
        "experience), loosely connected to the sermon's theme, e.g. \"Who is someone who kept loving you when you "
        "didn't deserve it?\" or \"What's something you've had to wait a long time for?\". Don't ask about God, "
        "the preacher or the passage yet. Note: one sentence on how it connects to the sermon.\n"
        "2-5. THE PASSAGE: four questions that follow the sermon in order, about specific words, verses, people "
        "or ideas in the text. Start from the best CANDIDATE QUESTIONS in the notes (improve the wording, keep the "
        "specific details and verse numbers of their answers). Make ONE of them about a QUOTE line from the notes, "
        f"copied word for word: \"{who} said, '<the quote>'. What did he mean by that?\". The four questions must "
        "each cover a different point. Notes: the answer as the sermon explained it, 1-3 sentences, with verse "
        "numbers.\n"
        "6. JESUS: how this passage points to Jesus and the gospel, as the sermon explained it.\n"
        "7. THIS WEEK: a personal application question based on what the sermon called people to do.\n"
        "Rules:\n"
        "- Exactly 7 questions, one per slot above, including the JESUS and THIS WEEK questions.\n"
        "- Every question gets 1-2 notes. A note gives the ANSWER; it must not repeat the question. Verse "
        "numbers go in the notes, not in the questions.\n"
        "- Use only facts, stories, quotes and verses from the notes above or the passage itself. Never "
        "invent them, and get right who a story is about. Only give verse numbers that appear in the notes; "
        "if unsure, leave the verse out.\n"
        f"- Plain, warm, conversational language. Questions start real discussion (not yes/no). Don't begin "
        f"more than two questions with \"{who}\".\n"
        'Reply as JSON: {"big_idea": "...", "main_points": ["..."], '
        '"questions": [{"q": "...", "notes": ["..."]}]}')


_REF = r"(?:[1-3] ?)?[A-Z][a-z]+\.? \d+(?::\d+(?:[-–]\d+)?)?"
_PAREN_REFS = re.compile(rf"\s*\((?:{_REF})(?:\s*[,;]\s*(?:{_REF}|\d+:\d+(?:[-–]\d+)?))*\)")


def _strip_refs(q):
    """The model likes to tack "(Romans 8:28-30)" onto questions; the references belong in the notes."""
    q = _PAREN_REFS.sub("", q).replace(" ?", "?").replace(" .", ".")
    return re.sub(r"([?!])(['’\"”])\.", r"\1\2", q).strip()


def generate(segments, info, status=lambda m: None, stop=lambda: False):
    """info: title, speaker, scripture, series, date (YYYY-MM-DD), start, end, examples (list of str).
    Returns the guide dict."""
    if not config.LLM_MODEL:
        raise RuntimeError("No AI model is set up (LLM_MODEL)")
    start, end = info.get("start") or 0, info.get("end") or 1e9
    words = " ".join(s["text"] for s in segments if start <= s["start"] <= end).split()
    if len(words) < 300:
        raise RuntimeError("The sermon transcript is too short to write a guide from")
    t0 = time.time()
    notes = _notes(words, info, status, stop)
    t1 = time.time()
    status("Writing the questions…")
    nt = _notes_text(notes)
    exs = info.get("examples") or []
    res = None
    for attempt in range(2):
        r = detect._ollama(_prompt(info, nt, exs), GUIDE_CTX, temperature=0.3, stop=stop,
                           options={"seed": random.randint(1, 1 << 30)})
        qs = [{"q": _strip_refs(str(q.get("q", ""))), "notes": [str(n).strip() for n in (q.get("notes") or []) if str(n).strip()]}
              for q in (r.get("questions") or []) if isinstance(q, dict) and str(q.get("q", "")).strip()]
        if len(qs) >= 5:
            res = r
            break
        status("Writing the questions again…")
    if not res:
        raise RuntimeError("The AI didn't write enough questions; try again")
    points = [str(p).strip().rstrip(".") for p in (res.get("main_points") or []) if str(p).strip()]
    if 2 <= len(points) <= 6:  # "Does anyone remember the main points?" (a staple of the guides), right after the opener
        who = _first_name(info.get("speaker")) or "The preacher"
        recall = {"q": f"{who} walked us through {len(points)} main points from {info.get('scripture') or 'the passage'}. "
                       "Does anyone remember what they were?",
                  "notes": [", ".join(f"{i}) {p}" for i, p in enumerate(points, 1)) + "."]}
        qs.insert(1 if len(qs) > 1 else len(qs), recall)
    return {
        "title": info.get("title") or "", "scripture": info.get("scripture") or "", "series": info.get("series") or "",
        "speaker": info.get("speaker") or "", "date": info.get("date") or "",
        "big_idea": str(res.get("big_idea") or "").strip(),
        "main_points": [str(p).strip() for p in (res.get("main_points") or []) if str(p).strip()],
        "questions": qs, "notes": notes,
        "meta": {"model": config.LLM_MODEL, "examples": len(exs), "words": len(words), "chunks": len(notes),
                 "seconds": round(time.time() - t0, 1), "notes_seconds": round(t1 - t0, 1), "at": time.time()},
    }


# --- formats ---------------------------------------------------------------------------------------
def _short_date(d):
    try:
        x = date.fromisoformat(d)
        return f"{x.month:02d}/{x.day}/{x.year % 100:02d}"
    except (TypeError, ValueError):
        return d or ""


def header_lines(g):
    title = f"“{g['title']}”" if g.get("title") else ""
    line2 = "—".join(x for x in (title, g.get("scripture")) if x)
    return [x for x in ("Life Group Questions", line2,
                        f"From the series “{g['series']}”" if g.get("series") else "", _short_date(g.get("date"))) if x]


def markdown(g):
    h = header_lines(g)
    out = [f"# {h[0]}", *[f"**{x}**  " if i == 0 else f"{x}  " for i, x in enumerate(h[1:])], ""]
    for i, q in enumerate(g["questions"], 1):
        out.append(f"{i}. {q['q']}")
        out.extend(f"    - *{n}*" for n in q["notes"])
    return "\n".join(out) + "\n"


def to_html(g):
    e = html.escape
    h = header_lines(g)
    parts = [f"<div class='gHead'><b>{e(h[0])}</b>" + "".join(f"<div>{e(x)}</div>" for x in h[1:]) + "</div><ol class='gQs'>"]
    for q in g["questions"]:
        parts.append(f"<li>{e(q['q'])}" + (("<ol type='a' class='gNotes'>" + "".join(f"<li>{e(n)}</li>" for n in q["notes"])
                                             + "</ol>") if q["notes"] else "") + "</li>")
    return "".join(parts) + "</ol>"


def filename(g):
    d = (g.get("date") or "").replace("-", "")
    name = f"{d} - {g.get('title') or 'Sermon'} Life Group Questions.docx"
    return re.sub(r'[\\/:*?"<>|]+', "", name).strip()


def to_docx(g, path: Path):
    """Word file in the style of a typical leader's guide: Times New Roman 12, a plain header, numbered
    questions with lettered leader notes in red underneath."""
    import docx
    from docx.enum.style import WD_STYLE_TYPE
    from docx.oxml import OxmlElement
    from docx.oxml.ns import qn
    from docx.shared import Pt, RGBColor, Inches

    d = docx.Document()
    sec = d.sections[0]
    sec.page_width, sec.page_height = Inches(8.5), Inches(11)
    sec.left_margin = sec.right_margin = sec.top_margin = sec.bottom_margin = Inches(1)
    normal = d.styles["Normal"]
    normal.font.name = "Times New Roman"
    normal.font.size = Pt(12)
    normal.element.rPr.rFonts.set(qn("w:eastAsia"), "Times New Roman")
    normal.paragraph_format.space_after = Pt(8)
    normal.paragraph_format.line_spacing = 1.25

    # numbering: 1. for questions, a. for the leader notes under each
    numbering = d.part.numbering_part.element
    abs_id = 90
    an = OxmlElement("w:abstractNum")
    an.set(qn("w:abstractNumId"), str(abs_id))
    mlt = OxmlElement("w:multiLevelType")
    mlt.set(qn("w:val"), "hybridMultilevel")
    an.append(mlt)
    for lvl, fmt, text, left in ((0, "decimal", "%1.", 360), (1, "lowerLetter", "%2.", 1080)):
        el = OxmlElement("w:lvl")
        el.set(qn("w:ilvl"), str(lvl))
        for tag, val in (("w:start", "1"), ("w:numFmt", fmt), ("w:lvlText", text), ("w:lvlJc", "left")):
            c = OxmlElement(tag)
            c.set(qn("w:val"), val)
            el.append(c)
        ppr = OxmlElement("w:pPr")
        ind = OxmlElement("w:ind")
        ind.set(qn("w:left"), str(left))
        ind.set(qn("w:hanging"), "360")
        ppr.append(ind)
        el.append(ppr)
        an.append(el)
    first_num = numbering.find(qn("w:num"))
    if first_num is not None:  # abstractNums must come before nums
        first_num.addprevious(an)
    else:
        numbering.append(an)
    num = OxmlElement("w:num")
    num_id = 90
    num.set(qn("w:numId"), str(num_id))
    ref = OxmlElement("w:abstractNumId")
    ref.set(qn("w:val"), str(abs_id))
    num.append(ref)
    numbering.append(num)

    def list_style(name, lvl, red):
        s = d.styles.add_style(name, WD_STYLE_TYPE.PARAGRAPH)
        s.base_style = normal
        s.paragraph_format.space_after = Pt(0 if lvl == 0 else 8)
        if lvl == 0:
            s.paragraph_format.space_before = Pt(6)
        if red:
            s.font.color.rgb = RGBColor(0xFF, 0, 0)
        ppr = s.element.get_or_add_pPr()
        npr = OxmlElement("w:numPr")
        il = OxmlElement("w:ilvl")
        il.set(qn("w:val"), str(lvl))
        ni = OxmlElement("w:numId")
        ni.set(qn("w:val"), str(num_id))
        npr.append(il)
        npr.append(ni)
        ppr.insert(0, npr)
        return s

    q_style, a_style = list_style("Question", 0, False), list_style("Answer", 1, True)
    for line in header_lines(g):
        p = d.add_paragraph(line)
        p.paragraph_format.space_after = Pt(0)
    d.add_paragraph().paragraph_format.space_after = Pt(0)
    for q in g["questions"]:
        d.add_paragraph(q["q"], style=q_style)
        for n in q["notes"]:
            d.add_paragraph(n, style=a_style)
    d.core_properties.title = " ".join(header_lines(g)[:2])
    d.core_properties.author = "Sermon Studio"
    tmp = path.with_name(f".{path.name}.part")
    d.save(str(tmp))
    tmp.replace(path)
    return path


# --- storage + background runs ------------------------------------------------------------------
def folder(jid):
    return DIR / jid


def load(jid):
    p = folder(jid) / "guide.json"
    return json.loads(p.read_text()) if p.exists() else None


def docx_path(jid):
    return folder(jid) / "guide.docx"


def save(jid, g):
    d = folder(jid)
    d.mkdir(parents=True, exist_ok=True)
    old = d / "guide.json"
    if old.exists():  # keep the previous version, in case the new one is worse
        old.replace(d / "guide.prev.json")
    (d / "guide.json").write_text(json.dumps(g, indent=1))
    (d / "guide.md").write_text(markdown(g))
    to_docx(g, docx_path(jid))


def running(jid):
    t = _running.get(jid)
    return bool(t and t.is_alive())


def start(job, segments, info, update, done=None):
    """Write the guide in a background thread. update(job, **fields) persists progress on the job
    (job["guide"]); done(job, guide) is called after it's saved."""
    with _lock:
        if running(job["id"]):
            raise RuntimeError("A guide is already being written for this sermon")

        def status(msg):
            update(job, guide={**(job.get("guide") or {}), "status": "working", "message": msg})

        def go():
            t0 = time.time()
            try:
                g = generate(segments, info, status)
                save(job["id"], g)
                update(job, guide={"status": "ready", "message": f"Written in {g['meta']['seconds']:.0f} s",
                                   "at": time.time(), "seconds": g["meta"]["seconds"], "examples": g["meta"]["examples"]})
                log.info("guide for %s: %d questions in %.0f s (%d words, %d examples)", job["id"],
                         len(g["questions"]), g["meta"]["seconds"], g["meta"]["words"], g["meta"]["examples"])
                if done:
                    done(job, g)
            except Exception as e:
                log.exception("guide for %s failed", job["id"])
                prev = load(job["id"])
                update(job, guide={"status": "ready" if prev else "error", "error": str(e),
                                   "message": f"Couldn't write the guide: {e}", "seconds": round(time.time() - t0, 1)})
        update(job, guide={"status": "working", "message": "Starting…", "started": time.time()})
        t = threading.Thread(target=go, daemon=True, name=f"guide-{job['id']}")
        _running[job["id"]] = t
        t.start()
